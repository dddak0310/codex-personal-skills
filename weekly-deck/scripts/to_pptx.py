#!/usr/bin/env python3
"""_export/slides/*.png + deck.json -> _export/*.pptx（每頁一張滿版圖，講稿進 speaker notes）。

⚠️ 截圖式 pptx 的文字在 PowerPoint 裡**不可編輯**。要改內容一律回頭改 deck.json
   → render_deck.py → shoot.py → 本腳本。備忘稿補回可搜尋的文字。

用法：
  to_pptx.py <deck.json> [-o out.pptx]
  to_pptx.py <deck.json> --slides <日期目錄>/_export/slides_zh -o <...>_zh.pptx

⚠️ 截圖與 pptx 都落在 `<日期目錄>/_export/` —— 那一層是**選用**的匯出品，
   不跑截圖／pptx 就不會存在（目錄形狀見 paths.py 尾的 WORK_STAGES）。
"""
import argparse, glob, json, os, sys

import paths

try:
    from pptx import Presentation
    from pptx.util import Inches, Emu
except ImportError:
    sys.exit("[pptx] 需要 python-pptx：python3 -m pip install --user python-pptx")

W_IN, H_IN = 13.333, 7.5   # 16:9


def notes_text(s, i):
    """備忘稿＝這頁的重點、承接／拋出、來源、以及日誌的執行細節。"""
    L = [f"[{s.get('id', f'S{i}')}] {s.get('title') or s.get('body', {}).get('title', '')}"]
    if s.get("point"):    L.append(f"重點：{s['point']}")
    if s.get("from"):     L.append(f"接續上一頁的結論：{s['from']}")
    if s.get("to"):       L.append(f"★ 這頁講完要說的結論：{s['to']}")
    if s.get("source"):   L.append(f"來源：{'、'.join(s['source'])}")
    if s.get("notes"):    L += ["", s["notes"]]
    return "\n".join(L)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("-o", "--out")
    ap.add_argument("--slides", help="截圖目錄（預設 <deck.json 同層>/_export/slides）")
    a = ap.parse_args()

    base = os.path.dirname(os.path.abspath(a.deck))
    deck = json.load(open(a.deck, encoding="utf-8"))
    slides = deck.get("slides", [])
    sdir = a.slides or os.path.join(paths.export_dir(base), "slides")
    pngs = sorted(glob.glob(os.path.join(sdir, "slide-*.png")))

    if len(pngs) != len(slides):
        sys.exit(f"[pptx] 圖片數({len(pngs)}) 與 deck.json 頁數({len(slides)}) 不符 —— "
                 f"請先重跑 shoot.py（不要帶 --only）")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W_IN), Inches(H_IN)
    blank = prs.slide_layouts[6]

    for i, (png, s) in enumerate(zip(pngs, slides), 1):
        sl = prs.slides.add_slide(blank)
        sl.shapes.add_picture(png, Emu(0), Emu(0),
                              width=Inches(W_IN), height=Inches(H_IN))
        sl.notes_slide.notes_text_frame.text = notes_text(s, i)

    meta = deck.get("meta", {})
    out = a.out or os.path.join(paths.export_dir(base, create=True),
                                f"weekly_report_{meta.get('week', 'deck')}.pptx")
    prs.save(out)
    print(f"[pptx] {len(pngs)} 頁 → {out}")


if __name__ == "__main__":
    main()
